"""API layer tests — team endpoints."""

from openpyxl import load_workbook


def _team(client):
    resp = client.get("/api/team")
    assert resp.status_code == 200
    return {m["name"]: m for m in resp.get_json()}


def test_list_team(flask_client):
    team = _team(flask_client)
    assert set(team) == {"Alice", "Bob"}
    assert team["Alice"]["start_year"] == 2025
    assert team["Alice"]["end_year"] is None


def test_add_member_creates_sheet(flask_client, patch_data_file):
    resp = flask_client.post("/api/team", json={"name": "Carol", "start_year": 2026})
    assert resp.status_code == 201
    assert resp.get_json() == {"name": "Carol", "start_year": 2026, "end_year": None}
    assert _team(flask_client)["Carol"]["start_year"] == 2026

    wb = load_workbook(patch_data_file)
    ws = wb["Carol 26"]
    assert ws["A2"].value == "Gennaio"
    assert ws["A14"].value == "Totale"
    assert ws["B14"].value == "=SUM(B2:B13)"


def test_add_existing_member_returns_400(flask_client):
    resp = flask_client.post("/api/team", json={"name": "Alice", "start_year": 2026})
    assert resp.status_code == 400
    assert "already exists" in resp.get_json()["error"]


def test_update_member_returns_current_info(flask_client):
    resp = flask_client.put("/api/team/Bob", json={})
    assert resp.status_code == 200
    assert resp.get_json()["name"] == "Bob"


def test_update_unknown_member_returns_404(flask_client):
    resp = flask_client.put("/api/team/Nobody", json={})
    assert resp.status_code == 404


def test_delete_member_keeps_sheets(flask_client, patch_data_file):
    resp = flask_client.delete("/api/team/Bob")
    assert resp.status_code == 204
    assert "Bob" in _team(flask_client)
    assert "Bob 26" in load_workbook(patch_data_file).sheetnames
